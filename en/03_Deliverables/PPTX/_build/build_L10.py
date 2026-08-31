"""
Build L10 EN PPTX: My Project, My Story (Graduation Day)
13 slides, Chaihuo brand rebuild from M0_EN_PPTPlan_CFG-5_Lesson10_v1.md
"""
import os
import sys
from pptx.util import Inches, Pt
from builder_lib import (
    new_deck, add_bg, add_rect, add_oval, add_text, add_multiline,
    add_title_bar, add_subtitle, add_footer, add_card, YELLOW, WHITE, INK,
    RED, CODEBG, PP_ALIGN, MSO_ANCHOR, set_notes
)

DECK_LABEL = "My Project, My Story"
TOTAL = 13
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..",
                   "M0_EN_Deck_CFG-5_Lesson10_MyProjectMyStory_v1.pptx")

COVER_ART = "/Users/leonfeng/Baiduyun/M0/M0-V2/素材/L10_封面插画_路演.png"


def page(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_footer(s, n, TOTAL, DECK_LABEL)
    return s


def num_block(s, x, y, num, size=0.34):
    add_rect(s, x, y, size, size, fill=YELLOW, border=None)
    add_text(s, x, y + 0.03, size, 0.32, num, size=15, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER)


def yellow_box(s, x, y, w, h, lines, size=13, line_spacing=1.25):
    add_rect(s, x, y, w, h, fill=YELLOW, border=None)
    if isinstance(lines, str):
        lines = [lines]
    add_multiline(s, x + 0.18, y + 0.12, w - 0.36, h - 0.24, lines, size=size,
                  line_spacing=line_spacing)


def grey_box(s, x, y, w, h, lines, size=10.5, line_spacing=1.35):
    add_card(s, x, y, w, h, fill=CODEBG, border=INK, border_w=1.0)
    items = [(t, {"mono": True}) if isinstance(t, str) else t for t in lines]
    add_multiline(s, x + 0.16, y + 0.10, w - 0.32, h - 0.20, items, size=size,
                  line_spacing=line_spacing)


def check_item(s, x, y, w, text, size=12, h=0.30):
    add_rect(s, x, y + 0.04, 0.20, 0.20, fill=YELLOW, border=None)
    add_text(s, x + 0.34, y, w - 0.34, h, text, size=size)


def red_line(s, x, y, w, h, title, body=None, body_size=13):
    add_rect(s, x, y, 0.07, h, fill=RED, border=None)
    add_oval(s, x + 0.18, y + 0.12, 0.11, fill=RED, border=None)
    add_text(s, x + 0.40, y + 0.06, w - 0.55, 0.35, title, size=18, bold=True)
    if body:
        add_text(s, x + 0.40, y + 0.42, w - 0.55, h - 0.50, body, size=body_size)


def add_photo_or_placeholder(s, path, x, y, w, h, label="photo"):
    if path and os.path.exists(path):
        s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        add_card(s, x, y, w, h, fill=WHITE, border=INK, border_w=1.5)
        add_text(s, x, y + h / 2 - 0.25, w, 0.50,
                 f"[ {label} — to be captured at rehearsal ]", size=11,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# ------------------------------------------------------------------ slides
def slide_01(prs):
    s = page(prs, 1)
    add_text(s, 0.62, 1.20, 6.4, 0.3,
             "Chaihuo Maker Academy · Smart Hardware Fundamentals (M0)",
             size=13)
    add_text(s, 0.62, 1.55, 8.76, 0.75, "My Project, My Story", size=38,
             bold=True)
    add_rect(s, 0.62, 2.50, 1.6, 0.06, fill=YELLOW, border=None)
    add_text(s, 0.62, 2.68, 6.4, 0.6,
             "The Final Leg & the Showcase · Graduation Day · "
             "Chaihuo Maker Academy · M0 Lesson 10", size=14, line_spacing=1.2)
    add_multiline(s, 0.62, 3.60, 5.8, 0.95, [
        ("Today: make your project understood — the release check, "
         "the README, the pitch, the roadshow —", {"size": 12.5}),
        ("and ten lessons, all lit.", {"size": 12.5}),
    ], line_spacing=1.35)
    add_photo_or_placeholder(s, COVER_ART, 6.80, 3.05, 2.45, 2.45,
                             label="roadshow illustration")
    set_notes(s, "Cover page only; students see it as they sit down. No "
                 "script. When the session starts, go straight to slide 02 — "
                 "the opener's 30-second recap is spoken, not projected. "
                 "Right side: the roadshow cover illustration reused from the "
                 "CN deck because it carries no Chinese glyphs. Footer: "
                 "Chaihuo Maker Academy · M0 · My Project, My Story | 01 / 13.")
    return s


def slide_02(prs):
    s = page(prs, 2)
    add_title_bar(s, "The Final-Day Map")
    add_rect(s, 0.62, 1.10, 0.07, 0.85, fill=YELLOW, border=None)
    add_text(s, 0.88, 1.12, 8.40, 0.85,
             "Last session: you called in reinforcements — code help and look help. "
             "Your project's function capped out, and it had a face for the first time. "
             "Baton handoff sheet? On the desk — we read it at the stand-up.",
             size=11.5)
    add_text(s, 0.62, 2.15, 8.76, 0.85,
             "THE LAST LESSON. GRADUATION DAY.", size=26, bold=True)
    add_text(s, 0.62, 3.05, 8.76, 0.55,
             "TODAY HAS EXACTLY ONE JOB — MAKE YOUR PROJECT UNDERSTOOD.",
             size=18, bold=True)
    add_multiline(s, 0.62, 3.80, 8.76, 0.95, [
        ("THE FINAL LEG — fix-only + the release check", {"size": 13, "bold": True}),
        ("> TELL THE WORLD — README, the math, the pitch, the rehearsal", {"size": 12.5}),
        ("> THE SHOWCASE & GRADUATION — the roadshow + the ceremony", {"size": 12.5}),
    ], line_spacing=1.40)
    add_text(s, 0.62, 4.95, 8.76, 0.25,
             "People who understand you are the ones your project deserves.",
             size=11.5, align=PP_ALIGN.CENTER)
    set_notes(s, "30-second recap: last session you called in reinforcements — "
                 "code help and look help. Your project's function capped out, "
                 "and it had a face for the first time. Baton handoff sheet? "
                 "On the desk — we read it at the stand-up. Today — the last "
                 "lesson. Graduation day. Three parts: the final leg, fix-only "
                 "→ tell the world (README, the math, the pitch) → the "
                 "showcase and graduation (the roadshow, the ceremony). Today "
                 "has exactly one job: make your project understood.")
    return s


def slide_03(prs):
    s = page(prs, 3)
    add_title_bar(s, "Leg-Three Stand-Up — Which Cut Do You Fix Today")
    add_card(s, 0.62, 1.50, 8.76, 0.80, fill=WHITE, border=INK, border_w=1.5)
    add_text(s, 0.86, 1.62, 8.28, 0.55,
             "READ YOUR BATON SHEET — what did last leg say is \"the first thing next leg\"?",
             size=13, bold=True)
    add_rect(s, 0.62, 2.50, 8.76, 0.78, fill=YELLOW, border=None)
    add_text(s, 0.80, 2.62, 8.40, 0.55,
             "\"I'M GOING TO FIX ___, SO THE DEMO HOLDS STEADY AT ___.\"",
             size=18, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_card(s, 0.62, 3.45, 8.76, 1.55, fill=WHITE, border=INK, border_w=1.5)
    add_text(s, 0.86, 3.58, 8.28, 1.35,
             "FIX-ONLY. The new feature you're itching to add? The \"later\" list welcomes it. "
             "A goal with \"add a ___\" in it gets stopped on the spot — that's an add. "
             "Make it a fix: which part are you least confident about? Fix that.",
             size=12.5, bold=True)
    set_notes(s, "Stand-up. First read your baton sheet — what did last leg "
                 "say is 'the first thing next leg'? Today, one sentence: "
                 "'I'm going to fix ___, so the demo holds steady at ___.' "
                 "Remember the discipline: fix-only. The new feature you're "
                 "itching to add? The 'later' list welcomes it.")
    return s


def slide_04(prs):
    s = page(prs, 4)
    add_title_bar(s, "Go — Two Jobs Only")
    add_multiline(s, 0.62, 1.45, 8.76, 1.00, [
        ("1) FIX THE FAULTS THE REVIEW LEFT BEHIND",
         {"size": 15, "bold": True}),
        ("2) POLISH DEMO STABILITY — FIVE CLEAN DEMOS IN A ROW, THAT'S \"STEADY\"",
         {"size": 15, "bold": True}),
    ], line_spacing=1.40)
    yellow_box(s, 0.62, 2.75, 8.76, 1.05, [
        "No shell yet? Don't start one today — if the Lesson-9 four moves can close it in 10 minutes, "
        "close it; otherwise it goes on stage as is. Function is the required question."],
        size=12, line_spacing=1.35)
    red_line(s, 0.62, 4.00, 8.76, 0.78,
             "FIX-ONLY",
             "new features go to the \"later\" list.")
    set_notes(s, "Go. Two jobs only. One: fix the faults the review left "
                 "behind. Two: polish demo stability — on roadshow day it must "
                 "work first try, in front of strangers. Five clean demos in a "
                 "row, that's 'steady.' No shell yet? Don't start one today — "
                 "if the Lesson-9 four moves can close it in 10 minutes, close "
                 "it; otherwise it goes on stage as is. Function is the "
                 "required question. When Act 1 ends, everyone stops — into "
                 "the release check.")
    return s


def slide_05(prs):
    s = page(prs, 5)
    add_title_bar(s, "Hands Off. The Release Check — Can It Face People?")
    add_text(s, 0.62, 1.42, 8.76, 0.95,
             "The review ran \"looks right\" → \"how sturdy is it?\" "
             "THIS time it asks \"can it face people?\" — the checklist goes "
             "line by line, and every line gets a conclusion.",
             size=12.5, bold=True, align=PP_ALIGN.CENTER)
    grey_box(s, 0.62, 2.50, 8.76, 1.85, [
        "Based on the current state: my project can ___ now, and can't ___ yet.",
        "You're a release inspector: in one hour it will demo in front of strangers.",
        "Give me a \"final pre-release checklist\": every item to confirm before the",
        "demo (enough power? cables seated? what's the first demo step?",
        "how do I save it if it fails?).",
    ], size=11, line_spacing=1.40)
    yellow_box(s, 0.62, 4.50, 8.76, 0.75, [
        "Tick every item. \"HOW DO I SAVE IT IF IT FAILS\" MUST HAVE AN ANSWER — "
        "none? Use the backup: the demo photo/video from your logs."],
        size=11.5, line_spacing=1.35)
    set_notes(s, "Hands off. The last review — and the persona upgrades: the "
                 "release check. Not fault-picking: 'can it face people?' "
                 "Have AI open the final checklist; this time you go through "
                 "it line by line, and every line gets a conclusion. Checklist "
                 "in hand, tick each item. The 'how do I save it if it fails' "
                 "line must have an answer — no answer? Use the backup: demo "
                 "photo/video from your logs.")
    return s


def slide_06(prs):
    s = page(prs, 6)
    add_title_bar(s, "The Engineer's Checklist — Brandy Ran These Four")
    checks = [
        ("DOES THE LIGHT COME ON IMMEDIATELY?", "the first sign of life"),
        ("DOES THE RECORDING END PROPERLY?", "the run is fully captured"),
        ("IS THE RECOGNIZED TEXT COMPLETE?", "no half-words or dropped lines"),
        ("DID THE PROMPT LAND IN THE RIGHT WINDOW?", "the project chat, not a fresh one"),
    ]
    positions = [(0.62, 1.50), (5.23, 1.50), (0.62, 2.65), (5.23, 2.65)]
    for i, ((x, y), (title, sub)) in enumerate(zip(positions, checks)):
        add_card(s, x, y, 4.15, 1.05, fill=WHITE, border=INK, border_w=1.5)
        num_block(s, x + 0.14, y + 0.14, str(i + 1), size=0.32)
        add_text(s, x + 0.56, y + 0.12, 3.35, 0.45, title, size=10.5, bold=True)
        add_text(s, x + 0.56, y + 0.62, 3.35, 0.40, sub, size=9.5)
    add_card(s, 0.62, 3.80, 8.76, 0.65, fill=WHITE, border=INK, border_w=1.5)
    add_rect(s, 0.62, 3.80, 0.07, 0.65, fill=YELLOW, border=None)
    add_text(s, 0.88, 3.92, 8.40, 0.45,
             "The first two ask: IS IT ACTUALLY WORKING? "
             "The last two ask: IS THE WORK COMPLETE?",
             size=12, bold=True)
    yellow_box(s, 0.62, 4.55, 8.76, 0.60, [
        "Map them onto YOUR project's equivalents and copy them into your checklist."],
        size=11)
    add_text(s, 0.62, 5.18, 8.76, 0.36,
             "Her fallback plan: serial logs, screen recording, a finished video — "
             "your \"how do I save it\" is exactly that.",
             size=9.5, align=PP_ALIGN.CENTER, line_spacing=1.20)
    set_notes(s, "Now the professional version — Brandy ran these four "
                 "before sharing her voice keyboard: does the light come on "
                 "immediately? does the recording end properly? is the "
                 "recognized text complete? did the prompt land in the right "
                 "window? The first two ask 'is it actually working?'; the "
                 "last two ask 'is the work complete?'. Map them onto your "
                 "project's equivalents and copy them into your checklist. "
                 "She also kept a fallback plan: serial logs, screen recording, "
                 "a finished video — your 'how do I save it' is exactly that.")
    return s


def slide_07(prs):
    s = page(prs, 7)
    add_title_bar(s, "README — Ten Logs Become One")
    add_text(s, 0.62, 1.42, 8.76, 0.55,
             "Your project's \"self-introduction + growth story\": someone who never sees the "
             "project, reading only this, understands what it is and how it got here.",
             size=11.5, bold=True, align=PP_ALIGN.CENTER)
    grey_box(s, 0.62, 2.08, 8.76, 1.05, [
        "These are my ten lesson logs: (paste all logs)",
        "Using only these logs, organize them into a README with this template:",
        "what it's called / whose problem it solves / how to use it (3 steps max) /",
        "how it grew (pick 3 turning points) / what's not perfect yet / what I'd still add.",
        "Do not add anything I didn't write.",
    ], size=10, line_spacing=1.30)
    grey_box(s, 0.62, 3.22, 8.76, 1.20, [
        "# My Project README",
        "## What it's called:",
        "## Whose problem it solves (a real person's name):",
        "## How to use it (3 steps max):",
        "## How it grew (3 turning points: the brief / the stuck / the fix):",
        "## What's not perfect yet (honestly, 1–2 lines):",
        "## What I'd still add (pick from the \"later\" list):",
    ], size=10, line_spacing=1.30)
    yellow_box(s, 0.62, 4.55, 8.76, 0.72, [
        "Generated? Line by line: DID IT INVENT ANYTHING? "
        "Caught one? Delete it — AI will confidently invent; proofreading is YOUR job. "
        "Ten lessons in, this is the final open-book exam."],
        size=10.5, line_spacing=1.30)
    set_notes(s, "Ten lessons, ten logs. Today they merge into a README — "
                 "your project's 'self-introduction + growth story': someone "
                 "who never sees the project, reading only this, understands "
                 "what it is and how it got here. Same old rule: AI organizes, "
                 "you proofread. After generation: line by line, did it "
                 "invent anything? Caught one? Delete it — AI will confidently "
                 "invent; proofreading is your job.")
    return s


def slide_08(prs):
    s = page(prs, 8)
    add_title_bar(s, "The Cost & Value Math — Count Exactly Two Things")
    grey_box(s, 0.62, 1.42, 8.76, 1.10, [
        "Count exactly two things:",
        "One. What it cost — the materials list (shell materials count, even a takeout box).",
        "Two. What it's worth — I think it's worth ___, because ___.",
        "(A high price is fine — but the reason must be \"it solves this problem,\"",
        "not \"I want it.\")",
    ], size=11, line_spacing=1.30)
    add_card(s, 0.62, 2.62, 4.20, 1.45, fill=WHITE, border=INK, border_w=1.5)
    add_text(s, 0.86, 2.78, 3.72, 0.35,
             "WHAT IT COST", size=13, bold=True)
    add_text(s, 0.86, 3.20, 3.72, 0.80,
             "the materials list", size=12, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    add_card(s, 5.18, 2.62, 4.20, 1.45, fill=YELLOW, border=INK, border_w=1.5)
    add_text(s, 5.42, 2.78, 3.72, 0.35,
             "WHAT IT'S WORTH", size=13, bold=True)
    add_text(s, 5.42, 3.20, 3.72, 0.80,
             "I think it's worth ___, because ___", size=11.5,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    yellow_box(s, 0.62, 4.30, 8.76, 0.78, [
        "The reason must be \"IT SOLVES THIS PROBLEM\" — not \"I want it.\" "
        "Can't make it hold? Lower the price until you can."],
        size=12, line_spacing=1.30)
    set_notes(s, "Rule first: today we count exactly two things — what it "
                 "cost, and what it's worth. No profits, no business talk. "
                 "Counting cost is applied math; saying value is an expression "
                 "exercise. Step one: list your materials — every component; "
                 "have AI check for gaps. Step two: if you sold it, what would "
                 "you price it at? And why? Write one number, write one "
                 "reason. The price can be above cost — even well above — but "
                 "the reason must be 'it solves this problem,' not 'I want it.'")
    return s


def slide_09(prs):
    s = page(prs, 9)
    add_title_bar(s, "The One-Page Pitch — Three Hard References")
    add_text(s, 0.62, 1.42, 8.76, 0.40,
             "One page that makes someone understand and care within 2 minutes.",
             size=12, bold=True, align=PP_ALIGN.CENTER)
    grey_box(s, 0.62, 1.90, 8.76, 1.85, [
        "# My One-Page Pitch",
        "",
        "## The problem (who is stuck with what — your requirements sheet's own words):",
        "",
        "## The solution (how my project fixes it + the demo move):",
        "",
        "## Who it's for (a real person from my Lesson-2 list):",
        "",
        "## What it's worth (quoting my account: it cost ___, I think it's worth ___, because ___)",
    ], size=10, line_spacing=1.30)
    add_multiline(s, 0.62, 3.92, 8.76, 1.05, [
        ("1) \"THE PROBLEM\" = YOUR REQUIREMENTS SHEET'S OWN WORDS",
         {"size": 11.5, "bold": True}),
        ("2) \"WHO IT'S FOR\" = A REAL PERSON FROM YOUR LESSON-2 LIST",
         {"size": 11.5, "bold": True}),
        ("3) \"WHAT IT'S WORTH\" = QUOTE YOUR ACCOUNT — ALL THREE REAL, AND THE PAGE HAS BONES",
         {"size": 11.5, "bold": True}),
    ], line_spacing=1.45)
    set_notes(s, "A pitch — one page — the page that makes someone understand "
                 "and care within 2 minutes. Four boxes: problem, solution, "
                 "who it's for, what it's worth. AI helps with structure and "
                 "wording — but three hard references it can't fake: 'who it's "
                 "for' must be a real person from your Lesson-2 list; 'what "
                 "it's worth' must quote the account you just did; 'the "
                 "problem' must be your requirements sheet's own words. All "
                 "three real, and the page has bones.")
    return s


def slide_10(prs):
    s = page(prs, 10)
    add_title_bar(s, "The AI Investor Rehearsal — Spend Your Fear Here")
    add_text(s, 0.62, 1.42, 8.76, 0.50,
             "Doesn't know tech · interrupts · only asks \"So what?\" and \"What's it to me?\"",
             size=11.5, bold=True, align=PP_ALIGN.CENTER)
    grey_box(s, 0.62, 2.05, 8.76, 1.55, [
        "You are an investor who understands nothing about technology.",
        "I'll read you my one-page pitch. You may interrupt me anytime,",
        "but you may only ask two kinds of questions: \"So what?\" and",
        "\"What's it to me?\" If you don't understand my answer,",
        "keep pressing until you do.",
    ], size=11, line_spacing=1.35)
    yellow_box(s, 0.62, 3.75, 8.76, 0.60, [
        "AI starts praising? Push back — \"No praise. Only 'so what?' questions.\""],
        size=13)
    add_text(s, 0.62, 4.55, 8.76, 0.55,
             "Every \"I can't answer\" is a gift — turn it into one line of your pitch, "
             "and nobody on stage can stump you.",
             size=12, align=PP_ALIGN.CENTER)
    set_notes(s, "The last gate before the stage: an investor who only "
                 "understands plain talk. Its persona: doesn't know tech, "
                 "interrupts, and only asks 'So what? What's it to me?' — "
                 "read your pitch to it; get stumped, revise, read again. "
                 "Getting stumped doesn't lose points — this is where you "
                 "spend your fear, so the stage has nothing left to take.")
    return s


def slide_11(prs):
    s = page(prs, 11)
    add_title_bar(s, "The Live Roadshow — Demo First, Then Talk")
    add_multiline(s, 0.62, 1.45, 8.76, 1.40, [
        ("1) DEMO FIRST, THEN TALK — THE THING REALLY MOVING SAYS MORE THAN ANYTHING",
         {"size": 16, "bold": True}),
        ("2) 2 MINUTES EACH — THE HOURGLASS EXEMPTS NOBODY",
         {"size": 16, "bold": True}),
        ("3) EVERY DESK ROTATES — EACH DESK SENDS ONE PERSON TO THE STAGE",
         {"size": 16, "bold": True}),
    ], line_spacing=1.50)
    yellow_box(s, 0.62, 2.95, 8.76, 0.85, [
        "Your job: 3 votes, cast after hearing everyone, "
        "for the project you'd MOST WANT TO SEE REALLY EXIST — you can't vote for yourself."],
        size=11.5, line_spacing=1.35)
    red_line(s, 0.62, 3.90, 8.76, 0.72,
             "THE HOURGLASS EXEMPTS NOBODY",
             "can't say it in 2 minutes = not clear yet. Finish that sentence, then wrap.")
    grey_box(s, 0.62, 4.70, 8.76, 0.55, [
        "ROADSHOW SELF-CHECK — tick before your turn:",
        "[ ] demo first: I can say my first demo step from memory",
        "[ ] five clean demos in a row   [ ] fits in 2 minutes (timed once)",
        "[ ] four pitch boxes, none longer than two sentences",
        "[ ] \"who it's for\" is a real person; \"what it's worth\" quotes my account",
        "[ ] survived one round of \"so what?\" questions",
        "[ ] backup in place: photo/video if the project dies",
        "[ ] I have an answer for \"how do I save it if it fails\"",
    ], size=8.5, line_spacing=1.20)
    set_notes(s, "Three rules: demo first, then talk — the thing really "
                 "moving says more than anything; 2 minutes each — the hourglass "
                 "exempts nobody (getting cut off is also the lesson: can't "
                 "say it in 2 minutes = not clear yet); every desk rotates — "
                 "and each desk sends one person to the stage. The audience has "
                 "a job too: 3 votes each, cast after hearing everyone, for "
                 "the project you'd most want to see really exist — you can't "
                 "vote for yourself.")
    return s


def slide_12(prs):
    s = page(prs, 12)
    add_title_bar(s, "Which One Would You Most Want to See Really Exist?")
    add_text(s, 0.62, 2.30, 8.76, 0.80,
             "3 VOTES — FOR OTHER PEOPLE'S PROJECTS. NOT YOURSELF.",
             size=24, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 0.62, 4.20, 8.76, 0.50,
             "Zero votes? The teacher names one specific highlight for every project — "
             "nobody walks away empty.",
             size=11.5, align=PP_ALIGN.CENTER)
    set_notes(s, "3 votes, for the project you'd most want to see really "
                 "exist — not yourself. After collecting, the teacher names "
                 "one specific highlight for every project — nobody walks away "
                 "empty. Voting is by bottle caps / sticky notes / show of "
                 "hands — no props bought.")
    return s


def slide_13(prs):
    s = page(prs, 13)
    add_title_bar(s, "The Closing Ceremony — Ten Lessons, All Lit")
    journey = [
        "Lesson 1 — you lit your first lamp",
        "Lesson 2 — you found a question worth doing",
        "Lesson 3 — your project grew a screen",
        "Lesson 4 — you got an AI team",
        "Lesson 5 — you taught hardware to see",
        "Lesson 6 — you moved your project home",
        "Lesson 7 — you briefed it",
        "Lesson 8 — you ran the first lap",
        "Lesson 9 — you called in reinforcements",
        "TODAY — YOU TOLD ITS STORY",
    ]
    jy = 1.42
    for i, line in enumerate(journey):
        sz = 11 if i < 9 else 13
        check_item(s, 0.62, jy, 8.76, line, size=sz, h=0.24)
        jy += 0.25
    yellow_box(s, 0.62, 4.00, 8.76, 0.45, [
        "You're not the one watching the magic anymore — you're the one making it."],
        size=11.5, line_spacing=1.20)
    add_text(s, 0.62, 4.45, 8.76, 0.28,
             "The \"did\" line — THE PROUDEST THING I MADE IN TEN LESSONS IS ___.",
             size=11, bold=True, align=PP_ALIGN.CENTER)
    add_rect(s, 0.62, 4.80, 8.76, 0.36, fill=RED, border=None)
    add_text(s, 0.78, 4.86, 8.44, 0.28,
             "THE CEREMONY BEGINS 15 MINUTES BEFORE THE END — NO SITUATION CUTS IT.",
             size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_rect(s, 0.62, 5.22, 8.76, 0.36, fill=RED, border=None)
    add_text(s, 0.78, 5.28, 8.44, 0.28,
             "THE PHOTO — HANDS ONLY, NEVER FACES. Hold up your project, hands up, backs to the camera.",
             size=10.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    set_notes(s, "Lesson one, you lit your first lamp; ... today — you told "
                 "its story. Ten lessons, all lit. You're the people who "
                 "walked the whole way. Remember Lesson 1? You lit your first "
                 "piece of hardware, your name was on the screen, and you said "
                 "'whoa.' Ten lessons later — your name is on your project, "
                 "in the README, in the applause just now. You're not the one "
                 "watching the magic anymore — you're the one making it. The "
                 "last log, four lines as usual. The 'did' line reads: the "
                 "proudest thing I made in ten lessons is ___. Photo — hold up "
                 "your project, hands up, backs to the camera; old rule, hands "
                 "only, never faces. That's a wrap!")
    return s


def main():
    prs = new_deck()
    slide_01(prs)
    slide_02(prs)
    slide_03(prs)
    slide_04(prs)
    slide_05(prs)
    slide_06(prs)
    slide_07(prs)
    slide_08(prs)
    slide_09(prs)
    slide_10(prs)
    slide_11(prs)
    slide_12(prs)
    slide_13(prs)
    prs.save(OUT)
    print(f"Saved: {OUT} slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()