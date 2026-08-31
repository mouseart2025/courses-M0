"""
Build the English Lesson 7 deck (Brief Your Project, 19 slides)
from M0_EN_PPTPlan_CFG-5_Lesson07_BriefYourProject_v1.md (brand rebuild).

Output: 交付物_EN/03_Deliverables/PPTX/M0_EN_Deck_CFG-5_Lesson07_BriefYourProject_v1.pptx
Layout per chaihuo-ppt-brand.md: white 70 / yellow 15 / ink 10 / red 5, 960x540.
Red budget (exactly 3): slide 05 (harassment boundary), slide 16 (the veto is
always yours), slide 17 (change the need, not the purchase). All other
warnings are bold black on yellow. No CJK / emoji / fullwidth glyphs on
screen; circled digits replaced with yellow number blocks; the fullwidth
vertical bar in the plan is rendered as an ASCII "|" or a middle dot.
No platform screenshots this lesson - pure text + line-drawn elements.
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from builder_lib import *
from pptx.util import Inches
from pptx.enum.text import MSO_ANCHOR

TOTAL = 19
DECK_LABEL = "Brief Your Project"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..",
                   "M0_EN_Deck_CFG-5_Lesson07_BriefYourProject_v1.pptx")


def page(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_footer(s, n, TOTAL, DECK_LABEL)
    return s


def num_block(s, x, y, num, size=0.40):
    add_rect(s, x, y, size, size, fill=YELLOW, border=None)
    add_text(s, x, y + 0.03, size, 0.32, num, size=15, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER)


def yellow_box(s, x, y, w, h, lines, size=13, line_spacing=1.25):
    add_rect(s, x, y, w, h, fill=YELLOW, border=None)
    add_multiline(s, x + 0.18, y + 0.12, w - 0.36, h - 0.24, lines, size=size,
                  line_spacing=line_spacing)


def grey_box(s, x, y, w, h, lines, size=10, line_spacing=1.3):
    add_card(s, x, y, w, h, fill=CODEBG, border=INK, border_w=1.0)
    items = [(t, {"mono": True}) if isinstance(t, str) else t for t in lines]
    add_multiline(s, x + 0.16, y + 0.10, w - 0.32, h - 0.20, items, size=size,
                  line_spacing=line_spacing)


def check_item(s, x, y, w, text, size=12, h=0.30):
    add_rect(s, x, y + 0.04, 0.20, 0.20, fill=YELLOW, border=None)
    add_text(s, x + 0.34, y, w - 0.34, h, text, size=size)


def red_line(s, y, text, size=13, h=0.30):
    add_rect(s, 0.62, y, 0.07, 0.46, fill=RED, border=None)
    add_oval(s, 0.86, y + 0.06, 0.12, fill=RED, border=None)
    add_text(s, 1.12, y + 0.02, 8.26, h, text, size=size, bold=True, color=RED)


def main():
    prs = new_deck()

    # ---------------------------------------------------------------- 01 cover
    s = page(prs, 1)
    add_text(s, 0.62, 1.20, 6.4, 0.3,
             "Chaihuo Maker Academy · Smart Hardware Fundamentals (M0)",
             size=13)
    add_text(s, 0.62, 1.55, 8.76, 0.75, "Brief Your Project", size=38,
             bold=True)
    add_rect(s, 0.62, 2.50, 1.6, 0.06, fill=YELLOW, border=None)
    add_text(s, 0.62, 2.68, 6.4, 0.6,
             "From the Project Wall to a Requirements Sheet · Chaihuo Maker "
             "Academy · M0 Lesson 7", size=14, line_spacing=1.2)
    add_text(s, 0.62, 3.45, 6.2, 0.35, "Lesson 7 · The milestone session",
             size=17, bold=True)
    add_multiline(s, 0.62, 3.92, 6.2, 0.95, [
        ("This morning: AI becomes the tester — your classmate finds fault in "
         "your project.", {"size": 12.5}),
        ("This afternoon: your big project gets its brief — and its building "
         "permit.", {"size": 12.5}),
    ], line_spacing=1.35)
    # sticky-note motif (line-drawn, 2 px black stroke)
    add_card(s, 7.30, 3.30, 1.80, 1.80, fill=YELLOW, border=INK,
             border_w=1.5)
    add_rect(s, 7.30, 4.80, 1.80, 0.30, fill=WHITE, border=INK, border_w=1.5)
    add_text(s, 7.30, 4.10, 1.80, 0.5, "the\nBrief Wall", size=11, bold=True,
             align=PP_ALIGN.CENTER, line_spacing=1.1)
    set_notes(s, "Cover page only; students see it as they sit down. No "
                 "script. When the session starts, go straight to slide 02 — "
                 "the opener's 30-second look-back is spoken, not projected. "
                 "Left side: a small sticky-note line icon (the Brief Wall "
                 "motif); yellow rule under the subtitle. Footer: Chaihuo "
                 "Maker Academy · M0 · Brief Your Project | 07 / 19.")

    # ---------------------------------------------------------------- 02 three things
    s = page(prs, 2)
    add_title_bar(s, "Today's Three Things")
    things = [
        ("1", "Finish last session's project — AI becomes the tester, your "
              "classmate tests your board, and we drag the faults out into "
              "the open", False),
        ("2", "THE BIG ONE: YOUR BIG PROJECT GETS ITS BRIEF TODAY — from the "
              "sentence you pinned on the wall in Lesson 2, all the way to "
              "\"ready to build\"", True),
        ("3", "Tour the parts pool, and stock up for your big project", False),
    ]
    ty = 1.15
    for num, txt, hi in things:
        if hi:
            add_rect(s, 0.62, ty - 0.06, 8.76, 0.78, fill=YELLOW, border=None)
        num_block(s, 0.80, ty, num)
        add_text(s, 1.42, ty - 0.01, 7.8, 0.72, txt, size=13,
                 line_spacing=1.2)
        ty += 0.92
    yellow_box(s, 0.62, 3.95, 8.76, 1.10, [
        "Lesson 5's link not finished? Report to a TA — during the "
        "self-check, five minutes to \"recognize scissors and light the "
        "LED,\" then join the peer test.",
        "Today isn't about how pretty the project is. It's about whether you "
        "can find fault."], size=12)
    set_notes(s, "30-second look-back: last session, you moved your project "
                 "into your own computer — from tenant to owner. Before that, "
                 "in Lesson 5, you taught your board to recognize rock, paper, "
                 "scissors — and to do something after recognizing: light up, "
                 "beep. And you found AI's biggest secret: it's not dumb — you "
                 "just taught it too little. Three things today. One: finish "
                 "last session's project — AI becomes the tester, your "
                 "classmate tests your board, and we drag the faults out into "
                 "the open. Two — the big one: your big project gets its brief "
                 "today — starting from the sentence you pinned on the wall in "
                 "Lesson 2, and walking it all the way to 'ready to build.' "
                 "Three: tour the parts pool, and stock up for your big "
                 "project. First thing first. Lesson 5's link not finished — "
                 "raise your hand. (note names) No problem: in a moment, "
                 "during the self-check, a TA runs you through the fallback "
                 "prompt — five minutes to 'recognize scissors and light the "
                 "LED.' Then you join the peer test normally. Still not "
                 "working? You share your neighbor's board and test as a team "
                 "— records count the same. A student missed Lesson 5 entirely "
                 "→ deploy the teacher's pre-trained model on a spare board; "
                 "they join the peer test straight away (the 'tester' role "
                 "doesn't need their own model). 'How far does the big project "
                 "have to go?' → 'The first five lessons were learning moves. "
                 "The marathon starts next session (Lesson 8) — this is what "
                 "you'll run. Today, you make it official.' The second thing "
                 "is the page's visual peak — the yellow highlight sells the "
                 "milestone.")

    # ---------------------------------------------------------------- 03 self-check
    s = page(prs, 3)
    add_title_bar(s, "Warm-Up Self-Check — Does It Still Know You?")
    add_text(s, 0.62, 1.02, 8.76, 0.30,
             "Plug in your little eye board — throw a scissors.", size=15,
             bold=True)
    add_card(s, 0.62, 1.50, 4.20, 1.75, fill=WHITE, border=INK, border_w=1.5)
    add_multiline(s, 0.80, 1.62, 3.84, 1.50, [
        ("KNOWS YOU", {"size": 13.5, "bold": True}),
        ("sit tight — next step", {"size": 12}),
    ], line_spacing=1.35)
    add_card(s, 5.18, 1.50, 4.20, 1.75, fill=WHITE, border=INK, border_w=1.5)
    add_multiline(s, 5.36, 1.62, 3.84, 1.50, [
        ("DOESN'T", {"size": 13.5, "bold": True}),
        ("think back to last session's attribution: did the light change? Did "
         "you move seats?", {"size": 10.5}),
        ("Quick rescue: against today's light, take 10 new photos and "
         "retrain. That's \"fix the data\" — you know how.",
         {"size": 10.5}),
    ], line_spacing=1.25)
    yellow_box(s, 0.62, 3.55, 8.76, 0.72, [
        "The board recognizes your gesture — or your retrain is underway."],
        size=13.5)
    set_notes(s, "Plug in your little eye board — throw a scissors. Does it "
                 "still know you? Knows you: sit tight. Doesn't: think back "
                 "to last session's attribution — did the light change? Did "
                 "you move seats? — Quick rescue: against today's light, take "
                 "10 new photos and retrain. That's 'fix the data' — you all "
                 "know how. Half the models dead (new room / new light — "
                 "common) → don't panic; this is the afternoon's first live "
                 "teaching material: 'See — new room, and it doesn't know you. "
                 "Guess what item #1 on today's trouble list is going to be. "
                 "You already know it.' Retraining over 8 minutes → stop; use "
                 "the 'accept' strategy: write 'this model only works at my "
                 "original seat' on the record, and enter the peer test. "
                 "Unfinished students catch up with TA B via the fallback "
                 "prompt in this window. Stand-still — the page is a "
                 "reference, not a lecture.")

    # ---------------------------------------------------------------- 04 trouble list
    s = page(prs, 4)
    add_title_bar(s, "Make AI Your Tester")
    add_text(s, 0.62, 1.00, 8.76, 0.26,
             "Remember Quinn — the one who helped you BREAK your Pomodoro "
             "timer in Lesson 4?", size=11.5)
    add_text(s, 0.62, 1.28, 8.76, 0.26,
             "This is Quinn's vision-edition moment. Send this — replace the "
             "brackets with your project:", size=11.5)
    grey_box(s, 0.62, 1.62, 8.76, 1.72, [
        "Quinn, my project: a board with a camera that can recognize "
        "scissors, rock, and paper.",
        "When it sees scissors it lights the LED; when it sees rock the "
        "buzzer beeps; when it",
        "can't recognize anything, everything stays off.",
        "You're a picky test engineer. Give me 5 tests for \"how to make it "
        "misjudge or fail,\"",
        "each one specific: what I should do, and what I expect to see.",
        "For example: hand gesture in backlight? Hand only half-visible? "
        "Switching gestures fast?",
    ], size=9.5, line_spacing=1.25)
    add_text(s, 0.62, 3.52, 8.76, 0.24, "FILTER RULES", size=10.5, bold=True)
    fr = [
        ("1", "Too vague (\"might be inaccurate\")? Send it back — ask for a "
              "specific scenario"),
        ("2", "Irrelevant to your project? Cross it off. KEEP AT LEAST 3 YOU "
              "CAN GENUINELY TRY — write them down."),
    ]
    fy = 3.82
    for num, txt in fr:
        num_block(s, 0.62, fy, num, size=0.34)
        add_text(s, 1.12, fy - 0.01, 8.26, 0.30, txt, size=11.5)
        fy += 0.44
    yellow_box(s, 0.62, 4.72, 8.76, 0.50, [
        "Your workbook holds a 3+ item list — each says \"how to do it, what "
        "you expect to see.\""], size=10.5)
    set_notes(s, "At the end of last session I said: today, AI becomes the "
                 "tester. Remember Quinn — the one who helped you break your "
                 "Pomodoro timer in Lesson 4? This is Quinn's vision-edition "
                 "moment. Send this to AI (project it), and replace the "
                 "brackets with your own project. Got your list? Don't test "
                 "yet — filter it first: items too vague ('might be "
                 "inaccurate') go back to AI for a specific scenario; items "
                 "irrelevant to your project get crossed off. Keep at least 3 "
                 "you can genuinely try. List too vague → the return prompt: "
                 "'Please rewrite item X as a concrete test: what action I "
                 "do, in what environment, what I watch for.' Time collapses "
                 "against the peer test → cut the list to protect the peer "
                 "test: 3 items are enough; the peer test's minutes are never "
                 "borrowed (it's this half's acceptance condition). "
                 "Stand-still. The prompt box is 6.1 verbatim — shared text "
                 "with the Teacher's Guide and the workbook; any edit must "
                 "sync all three.")

    # ---------------------------------------------------------------- 05 peer test (RED 1)
    s = page(prs, 5)
    add_title_bar(s, "Swap Projects — Test Item by Item")
    add_multiline(s, 0.62, 1.05, 8.76, 0.80, [
        ("In your hands: your neighbor's board + the trouble list they "
         "wrote.", {"size": 14, "bold": True}),
        ("Test item by item — record on their list: pass, checkmark; broke, "
         "write exactly how it broke. Then they test yours. BOTH LISTS GET "
         "FILLED.", {"size": 12}),
    ], line_spacing=1.3)
    add_text(s, 0.62, 2.02, 8.76, 0.24, "THREE RULES", size=10.5, bold=True)
    num_block(s, 0.62, 2.32, "1", size=0.34)
    add_text(s, 1.12, 2.30, 8.26, 0.56,
             "Find a fault? Don't laugh — the list you hand over is a knife, "
             "and a badge of honor too: your list getting tested to the "
             "breaking point means your neighbor is a real professional",
             size=10.5, line_spacing=1.2)
    num_block(s, 0.62, 2.94, "2", size=0.34)
    add_text(s, 1.12, 2.92, 8.26, 0.30,
             "After you test theirs, they test yours — both lists get filled",
             size=11.5)
    num_block(s, 0.62, 3.36, "3", size=0.34)
    add_rect(s, 0.62, 3.34, 0.07, 0.68, fill=RED, border=None)
    add_oval(s, 0.86, 3.40, 0.12, fill=RED, border=None)
    add_multiline(s, 1.12, 3.36, 8.26, 0.64, [
        ("THE HARASSMENT IS ONLY \"SHOW IT THINGS\" —",
         {"size": 11.5, "bold": True, "color": RED}),
        ("NO TOUCHING WIRES, NO TAKING PARTS OFF, NO COVERING ANYTHING "
         "EXCEPT THE LENS",
         {"size": 11.5, "bold": True, "color": RED}),
    ], line_spacing=1.2)
    yellow_box(s, 0.62, 4.10, 8.76, 0.60, [
        "Both lists carry real results — checkmarks and \"here's how it "
        "broke.\""], size=13)
    set_notes(s, "Swap projects. In your hands: your neighbor's board + the "
                 "trouble list they wrote. Test item by item, record on their "
                 "list: pass, checkmark; broke, write exactly how it broke. "
                 "Three rules. One: the harassment is only 'show it things' — "
                 "no touching wires, no taking parts off, no covering "
                 "anything except the lens. Two: if you find a fault, don't "
                 "laugh at them — the list you hand over is a knife, and it's "
                 "a badge of honor too: your list getting tested to the "
                 "breaking point means your neighbor is a real professional. "
                 "Three: after you test theirs, they test yours — both lists "
                 "get filled. The red line is the deck's first red — the "
                 "hardware boundary. The peer test turns into actual sabotage "
                 "(pulling wires / covering the lens) → restate the rules "
                 "once; repeat offenders get promoted to power-cycle tester "
                 "(their official test: 'unplug the power and plug it back — "
                 "does it recover?' — which is, genuinely, an effective "
                 "test). One side's project doesn't run → pair up on the "
                 "working one; the student without a running board takes the "
                 "tester+recorder role (records are still a deliverable). "
                 "Finished early → free-form escalation: 'Forget the list. "
                 "Think of the nastiest move yourself.' This block is this "
                 "half's acceptance condition — never compress it. "
                 "Stand-still.")

    # ---------------------------------------------------------------- 06 accept or reject
    s = page(prs, 6)
    add_title_bar(s, "Accept or Reject — One Fix Round")
    add_multiline(s, 0.62, 1.02, 8.76, 0.56, [
        ("Take back your own list — look at what it's become: your project's "
         "full list of sins.", {"size": 12.5}),
        ("Judge them one by one: ACCEPT, or REJECT?", {"size": 13.5,
                                                       "bold": True}),
    ], line_spacing=1.25)
    add_text(s, 0.62, 1.72, 4.30, 0.24, "REJECT → FIX IT NOW", size=11,
             bold=True)
    fixes = [
        ("FIX THE DATA", "back into SenseCraft, photograph the scene where "
                         "it broke"),
        ("FIX THE LOGIC", "have Codecraft change a condition, like \"only "
                          "counts if it recognizes it twice in a row\""),
        ("FIX THE THRESHOLD", "raise the confidence line from 70% to 85%"),
    ]
    fy = 2.02
    for head, body in fixes:
        add_card(s, 0.62, fy, 4.30, 0.78, fill=WHITE, border=INK,
                 border_w=1.5)
        add_multiline(s, 0.78, fy + 0.08, 3.98, 0.64, [
            (head, {"size": 11.5, "bold": True}),
            (body, {"size": 9.5}),
        ], line_spacing=1.2)
        fy += 0.90
    add_text(s, 5.20, 1.72, 4.18, 0.24, "ACCEPT → WRITE ONE LINE", size=11,
             bold=True)
    add_rect(s, 5.20, 2.02, 4.18, 1.42, fill=YELLOW, border=None)
    add_multiline(s, 5.38, 2.14, 3.82, 1.20, [
        ("\"I know it fails when ___, and I accept it.\"", {"size": 12.5,
                                                            "bold": True}),
        ("You made this call in Lesson 2. Today you make it more "
         "professionally.", {"size": 10.5}),
    ], line_spacing=1.3)
    add_rect(s, 5.20, 3.62, 4.18, 1.48, fill=WHITE, border=INK, border_w=1.5)
    add_rect(s, 5.20, 3.62, 0.07, 1.48, fill=YELLOW, border=None)
    add_multiline(s, 5.44, 3.72, 3.80, 1.30, [
        ("Not sure which path? Send the failure to AI:", {"size": 10.5}),
        ("\"My model misjudges in ___ situation. Fix the data, fix the "
         "logic, or fix the threshold — which do you recommend, and why?\"",
         {"size": 10}),
        ("Advice is advice. WHICH PATH IS STILL YOUR CALL.",
         {"size": 10.5, "bold": True}),
    ], line_spacing=1.25)
    set_notes(s, "Take back your own list — look at what it's become: your "
                 "project's full list of sins. Now judge them one by one: "
                 "Accept or Reject. Reject? Fix it now. Three paths (write on "
                 "the board): fix the data — go back into SenseCraft and "
                 "photograph the scene where it broke; fix the logic — have "
                 "Codecraft change a condition, like 'only counts if it "
                 "recognizes it twice in a row'; fix the threshold — raise "
                 "the confidence line from 70% to 85%. Not sure which path? "
                 "Don't wait for me — send the failure to AI: 'My model "
                 "misjudges in ___ situation. Fix the data, fix the logic, or "
                 "fix the threshold — which do you recommend, and why?' "
                 "Advice is advice — which path is still your call. Fixing "
                 "the logic and threshold also goes to AI; your only job is "
                 "the old rule: say the input and output clearly. Accept? "
                 "Write one line: 'I know it fails when ___, and I accept "
                 "it.' You made this call in Lesson 2 — today you make it "
                 "more professionally. Everyone chooses 'accept' (too lazy to "
                 "fix) → require at least one rejection: 'Pick the fault "
                 "you'd most hate a judge to catch live.' Everyone chooses "
                 "'reject' (perfectionism) → remind them of the time: 'You "
                 "can't fix 5 faults in 15 minutes. Pick the most deadly one "
                 "and fix it; the rest get accepted — that's what a "
                 "trade-off looks like.' Someone fixes via 'the data' → "
                 "point at the loop closing: 'See — you're back at last "
                 "session's attribution: whatever situation breaks it is "
                 "whatever data you never fed it.' Stand-still.")

    # ---------------------------------------------------------------- 07 share-out
    s = page(prs, 7)
    add_title_bar(s, "Share-Out + Log Line One")
    add_multiline(s, 0.62, 1.15, 8.76, 0.90, [
        ("One pick per table: \"the most vicious misread + my counter\"",
         {"size": 18, "bold": True}),
        ("30 seconds to the class.", {"size": 13}),
    ], line_spacing=1.35)
    add_card(s, 0.62, 2.45, 8.76, 1.90, fill=WHITE, border=INK, border_w=1.5)
    add_multiline(s, 0.84, 2.60, 8.32, 1.62, [
        ("Four lines as usual — today's \"done\" line:",
         {"size": 12, "bold": True}),
        ("I HAD AI OPEN A TROUBLE LIST, MY CLASSMATE FOUND ___ FAULTS IN MY "
         "PROJECT, I FIXED ___ AND ACCEPTED ___", {"size": 14, "bold": True}),
        ("(fill the blanks in your workbook — the yellow slots)",
         {"size": 10.5}),
    ], line_spacing=1.6)
    set_notes(s, "One pick per table: 'the most vicious misread + my counter' "
                 "— 30 seconds to the class. Then log line one. Four lines as "
                 "usual; today 'done' reads: I had AI open a trouble list, my "
                 "classmate found ___ faults in my project, I fixed ___ and "
                 "accepted ___. Collect the best one-liners — they're "
                 "showcase material (the reflection page asks for them). Fast "
                 "flip; no stand-still needed — the log card stays up while "
                 "students write.")

    # ---------------------------------------------------------------- 08 Neil's line
    s = page(prs, 8)
    add_title_bar(s, "Neil's Line")
    add_subtitle(s, "first, listen to someone for a minute — Neil, founder "
                    "of MIT's Fab Academy", y=1.00, size=12)
    add_rect(s, 0.62, 1.50, 0.09, 1.55, fill=YELLOW, border=None)
    add_multiline(s, 0.94, 1.55, 8.30, 1.48, [
        ("Learning to make things is not about making what you can BUY IN A "
         "STORE — that's not yours to make.", {"size": 14}),
        ("It's about making what you CAN'T BUY — personalized, for yourself "
         "and the people you care about.", {"size": 14}),
    ], line_spacing=1.3)
    add_rect(s, 0.62, 3.20, 0.09, 0.95, fill=YELLOW, border=None)
    add_multiline(s, 0.94, 3.24, 8.30, 0.90, [
        ("And a project idea isn't a \"good idea\" you think up — IT GROWS "
         "OUT OF YOUR OWN EXPERIENCE — the things that annoy you, the things "
         "the people you care about put up with every day.",
         {"size": 13}),
    ], line_spacing=1.25)
    yellow_box(s, 0.62, 4.32, 8.76, 0.90, [
        "A flashy idea copied from the internet won't survive three lessons. "
        "A problem that grew out of your experience can survive four "
        "marathons. Fab Academy students write down their Final Project in "
        "week one — and for twenty weeks, everything they do is collecting "
        "modules for it. That's exactly what you're doing today."], size=10.5)
    set_notes(s, "Second half. First, listen to someone for a minute — Neil, "
                 "the founder of MIT's Fab Academy. He says: learning to make "
                 "things is not about making what you can buy in a store — "
                 "that's not yours to make. It's about making what you can't "
                 "buy — personalized, for yourself and the people you care "
                 "about. And he keeps telling students: a project idea isn't a "
                 "'good idea' you think up — it grows out of your own "
                 "experience — the things that annoy you, the things the "
                 "people you care about put up with every day. A flashy idea "
                 "copied from the internet won't survive three lessons. A "
                 "problem that grew out of your experience can survive four "
                 "marathons. Fab Academy students write down their Final "
                 "Project in week one — and for twenty weeks, everything they "
                 "do is collecting modules for it. That's exactly what you're "
                 "doing today. First mention of Neil — introduce him in one "
                 "breath ('MIT's Fab Academy'), no school history. This page "
                 "is one of the two golden-line peaks of the brief arc (the "
                 "other is slide 18) — big type, wide white space, no task "
                 "list. A NEW page the CN v2 plan never had.")

    # ---------------------------------------------------------------- 09 directions
    s = page(prs, 9)
    add_title_bar(s, "Same Skills — What Did They Build?")
    add_multiline(s, 0.62, 1.00, 8.76, 0.54, [
        ("What Neil's course actually teaches: sensors, circuits, embedded "
         "systems, structural design — the same family as what's in your "
         "hands.", {"size": 11}),
        ("Look at what past students built with this skill set:",
         {"size": 11}),
    ], line_spacing=1.25)
    dirs = [
        "protective gear (a lab made it during the pandemic)",
        "homemade oscilloscopes", "homemade microscopes",
        "a 10-km link out of a parabolic antenna (in the mountains)",
        "educational satellites", "soft robots", "drones",
        "hydroponic systems & homemade instruments",
    ]
    for i, d in enumerate(dirs):
        cx = 0.62 + (i % 4) * 2.22
        cy_ = 1.68 + (i // 4) * 1.20
        add_card(s, cx, cy_, 2.10, 1.08, fill=WHITE, border=INK,
                 border_w=1.5)
        add_text(s, cx + 0.10, cy_ + 0.08, 1.90, 0.92, d, size=9.5,
                 anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.2)
    add_rect(s, 0.62, 4.24, 0.07, 0.92, fill=YELLOW, border=None)
    add_multiline(s, 0.88, 4.28, 8.50, 0.86, [
        ("NOT ONE OF THEM IS A FLASHY IDEA COPIED FROM THE INTERNET — EVERY "
         "ONE IS A REAL PROBLEM FROM THE MAKER'S OWN EXPERIENCE.",
         {"size": 11.5, "bold": True}),
        ("Your topic can grow from these directions too.", {"size": 11}),
    ], line_spacing=1.25)
    set_notes(s, "What does Neil's course actually teach? Sensors, circuits, "
                 "embedded systems, structural design — the same family as "
                 "what's in your hands. Look at what past students built with "
                 "this skill set (scan down the list): during the pandemic, a "
                 "lab making protective gear; homemade oscilloscopes and "
                 "microscopes; someone in the mountains pulling a "
                 "10-kilometer link out of a parabolic antenna; educational "
                 "satellites, soft robots, drones, hydroponic systems, "
                 "homemade instruments. Notice: not one of them is a flashy "
                 "idea copied from the internet — every one is a real problem "
                 "from the maker's own experience. Your topic can grow from "
                 "these directions too. Source: the Fab Academy graduation "
                 "lecture (vimeo 949353). The icons degrade to plain text "
                 "cards (names only — the page still lands). A NEW page the "
                 "CN v2 plan never had.")

    # ---------------------------------------------------------------- 10 four notes
    s = page(prs, 10)
    add_title_bar(s, "Neil's Four Notes")
    add_subtitle(s, "what counts as a good project?", y=1.00, size=13)
    notes4 = [
        ("A MASTERPIECE MEANS SKILLED, NOT GRAND",
         "a finished ordinary project beats an unfinished grand vision"),
        ("SPIRAL DEVELOPMENT", "every lap ends with a running version"),
        ("MODULAR", "cut it into small pieces, build and test each "
                    "separately"),
        ("DOCUMENT AS YOU GO", "the documentation grows with the project"),
    ]
    for i, (head, body) in enumerate(notes4):
        cx = 0.62 + (i % 2) * 4.48
        cy_ = 1.45 + (i // 2) * 1.30
        add_card(s, cx, cy_, 4.28, 1.16, fill=WHITE, border=INK,
                 border_w=1.5)
        num_block(s, cx + 0.16, cy_ + 0.14, str(i + 1), size=0.34)
        add_text(s, cx + 0.64, cy_ + 0.13, 3.50, 0.50, head, size=11.5,
                 bold=True, line_spacing=1.1)
        add_text(s, cx + 0.64, cy_ + 0.66, 3.50, 0.44, body, size=10.5,
                 line_spacing=1.2)
    yellow_box(s, 0.62, 4.12, 8.76, 1.05, [
        "Sound familiar? You already know all four — Lesson 4's Pomodoro "
        "timer was built lap by lap in a spiral; your four-line log IS "
        "documentation.",
        "Today's requirements sheet is the first line of your project's "
        "document."], size=11.5)
    set_notes(s, "So what counts as a good project? Neil's four notes. One: a "
                 "masterpiece means skilled, not grand — a finished ordinary "
                 "project beats an unfinished grand vision. Two: spiral "
                 "development — every lap ends with a running version. "
                 "Three: modular — cut it into small pieces, build and test "
                 "each separately. Four: document as you go — the "
                 "documentation grows with the project. Sound familiar? You "
                 "already know all four: Lesson 4's Pomodoro timer was built "
                 "lap by lap in a spiral; your four-line log is documentation. "
                 "Today's requirements sheet is the first line of your "
                 "project's document. The 'you already know all four' line is "
                 "the connective tissue — it ties the four notes to the "
                 "students' own experience. A NEW page the CN v2 plan never "
                 "had.")

    # ---------------------------------------------------------------- 11 project wall
    s = page(prs, 11)
    add_title_bar(s, "The Project Wall — Do You Still Love It?")
    add_multiline(s, 0.62, 1.00, 8.76, 0.52, [
        ("In Lesson 2, you pinned a sentence here: I'm making a ___ for ___ "
         "because ___.", {"size": 11.5}),
        ("Four weeks later — now look at your own sentence.", {"size": 11.5}),
    ], line_spacing=1.25)
    add_text(s, 0.62, 1.58, 8.76, 0.42, "DO YOU STILL LOVE THIS TOPIC?",
             size=24, bold=True, align=PP_ALIGN.CENTER)
    add_card(s, 0.62, 2.10, 4.28, 0.78, fill=WHITE, border=INK, border_w=1.5)
    add_multiline(s, 0.80, 2.20, 3.92, 0.60, [
        ("LOVE IT", {"size": 12, "bold": True}),
        ("sit back down; we'll deepen it in a minute", {"size": 10.5}),
    ], line_spacing=1.2)
    add_card(s, 5.10, 2.10, 4.28, 0.78, fill=WHITE, border=INK, border_w=1.5)
    add_multiline(s, 5.28, 2.20, 3.92, 0.60, [
        ("DON'T LOVE IT / UNSURE", {"size": 12, "bold": True}),
        ("you can change it. Changing today is not embarrassing. Carrying "
         "something you don't love all the way to Lesson 10 — that's the "
         "embarrassing one.", {"size": 9}),
    ], line_spacing=1.15)
    add_text(s, 0.62, 3.02, 8.76, 0.24,
             "RE-TOPIC SELF-CHECK — 5-MINUTE TIMER", size=10.5, bold=True)
    checks3 = [
        (3.28, 0.30, "Do I really not love it, or am I just stuck?"),
        (3.64, 0.46, "Does the new topic come from my own experience — who's "
                     "it for? Can I name a real person and one real time they "
                     "were annoyed?"),
        (4.14, 0.30, "Can the boards in my hands actually make it?"),
    ]
    for i, (cy_, ch, c) in enumerate(checks3):
        num_block(s, 0.62, cy_, str(i + 1), size=0.30)
        add_text(s, 1.06, cy_ - 0.01, 8.32, ch, c, size=10.5)
    yellow_box(s, 0.62, 4.50, 8.76, 0.72, [
        "All three answered: change. Can't answer: stay on the original — "
        "stuck? raise your hand, a TA gets you past it. You can't change to "
        "a \"flashy copied idea\" — change only to a real problem from your "
        "own experience."], size=10)
    set_notes(s, "Everyone up — walk to the Project Wall. In Lesson 2, you "
                 "pinned a sentence here: I'm making a ___ for ___ because "
                 "___. Four weeks later — you've lit up hardware, led an AI "
                 "team, taught a board to recognize things. Now look at your "
                 "own sentence, and I ask you: do you still love this topic? "
                 "Love it: sit back down; we'll deepen it in a minute. Don't "
                 "love it — or you're unsure: you can change it. Changing "
                 "today is not embarrassing. Carrying something you don't "
                 "love all the way to Lesson 10 — that's the embarrassing "
                 "one. Changing rules: 5 minutes, a three-question self-check "
                 "(projected): all three answered: change. Can't answer: stay "
                 "on the original, and if you're stuck, raise your hand and a "
                 "TA gets you past it. Note rule two: you can't change to a "
                 "'flashy copied idea' — change only to a real problem from "
                 "your own experience. Still can't decide in 5 minutes — "
                 "come to me. I have three backup topics ready; done "
                 "beautifully, they're still good projects. Nothing to be "
                 "ashamed of. The teacher's hand itches to pick a topic for a "
                 "student (this lesson's biggest teaching accident): "
                 "self-check your language — you may only ask 'Who did you "
                 "say it's for? When does that person hit this problem?' — "
                 "never 'I think you should make ___ instead.' However "
                 "ordinary the answer, it's the student's topic. Re-topic "
                 "rate over 30% → don't handle it live; note it on the "
                 "reflection page (the Lesson 2 → today gap needs review). A "
                 "student's topic collapses and they're genuinely lost → give "
                 "a backup topic verbally: 'Take one and start building; your "
                 "own idea will come while you work — the requirements sheet "
                 "is allowed to change. Cross it out; don't erase it.' This "
                 "page is the slimmed version of the CN plan's page 08.")

    # ---------------------------------------------------------------- 12 PM talk
    s = page(prs, 12)
    add_title_bar(s, "Summon Your Product Manager — John")
    add_subtitle(s, "go deep", y=1.00, size=13)
    add_rect(s, 0.62, 1.38, 8.76, 0.94, fill=YELLOW, border=None)
    add_multiline(s, 0.80, 1.46, 8.40, 0.80, [
        ("The opener MUST carry your Lesson 2 real user's name — a need "
         "without a real name is a fake need.", {"size": 11.5, "bold": True}),
        ("Watch John's first question: \"When did this last happen?\" — can "
         "you answer with a specific day? Then it's a real experience. "
         "Can't? The topic is still just an idea.", {"size": 10.5}),
    ], line_spacing=1.25)
    grey_box(s, 0.62, 2.44, 8.76, 1.86, [
        "John, my topic is \"I'm making a (what) for (a real person's name),",
        "because (the real problem they have).\"",
        "You're my product manager. Interview me and go deep:",
        "1. In what situation does this problem happen? When did it last "
        "happen?",
        "2. Without this thing, how do they get by right now?",
        "3. When and where will they use the thing I make?",
        "When you're done asking, organize my answers into three sentences.",
    ], size=10, line_spacing=1.3)
    add_rect(s, 0.62, 4.36, 0.07, 0.32, fill=YELLOW, border=None)
    add_text(s, 0.88, 4.38, 8.50, 0.30,
             "Keep going in the conversation you already have — no new "
             "windows. Open a new one and every conclusion so far is gone.",
             size=10.5, bold=True)
    yellow_box(s, 0.62, 4.74, 8.76, 0.48, [
        "AI's three sentences contain a real name and a real scene — copied "
        "into your workbook."], size=9.5)
    set_notes(s, "Back to your seats — summon your product manager. John — "
                 "you know this role. Today we go deep. The opener (project "
                 "it), and note: it must carry your Lesson 2 real user's name "
                 "— a need without a real name is a fake need. How does "
                 "Neil's 'start from personal experience' get tested here? "
                 "Watch John's first question: 'When did this last happen?' — "
                 "if you can answer with a specific day, it's a real "
                 "experience; if you can't, the topic is still just an idea. "
                 "Old discipline: keep going in the conversation you already "
                 "have — no new windows. Open a new one and every conclusion "
                 "so far is gone. User written as 'everyone / people' → TA A "
                 "sends it back: 'Say one person's name. Who was on your "
                 "Lesson 2 list?' The chat contains none of the student's own "
                 "information (AI inventing needs on its own) → the "
                 "circulation red line — the citation check — send it back to "
                 "fill in. A student actually interviewed someone (asked "
                 "mom/sister over the weekend) → invite a 30-second share; "
                 "set the benchmark. Output anchor 5: the PM's three "
                 "sentences. Stand-still. The prompt box is 6.2 verbatim — "
                 "shared text; any edit must sync all three documents.")

    # ---------------------------------------------------------------- 13 UX talk
    s = page(prs, 13)
    add_title_bar(s, "Summon the Experience Designer — Sally")
    add_subtitle(s, "one thing: the usage scenario", y=1.00, size=13)
    add_card(s, 0.62, 1.42, 8.76, 1.12, fill=WHITE, border=INK, border_w=1.5)
    add_rect(s, 0.62, 1.42, 0.07, 1.12, fill=YELLOW, border=None)
    add_multiline(s, 0.88, 1.52, 8.36, 0.94, [
        ("\"Mom walks past the coffee table; it's been two hours since the "
         "water cup moved, so it flashes three times and gives one soft "
         "ding.\"", {"size": 11.5}),
        ("A PERSON, A PLACE, AN ACTION, A REACTION. Write yours in that "
         "shape.", {"size": 11.5, "bold": True}),
    ], line_spacing=1.3)
    grey_box(s, 0.62, 2.68, 8.76, 1.72, [
        "Sally, based on the previous round: (paste John's three sentences).",
        "You're my experience designer. Write me a \"usage scenario\":",
        "who uses it, where, how they use it, and how it reacts. Use a "
        "specific",
        "person and place, no more than four sentences. Then tell me:",
        "what does this thing look like, and where does it live?",
    ], size=10, line_spacing=1.3)
    yellow_box(s, 0.62, 4.54, 8.76, 0.68, [
        "Your scenario can point to all four — who, where, how, reaction — "
        "copied into your workbook or notebook."], size=11)
    set_notes(s, "Summon the experience designer. Sally. She produces one "
                 "thing: the usage scenario — who uses it, where, how. Watch "
                 "my example (project it): 'Mom walks past the coffee table; "
                 "it's been two hours since the water cup moved, so it "
                 "flashes three times and gives one soft ding.' — a person, "
                 "a place, an action, a reaction. Write yours in that shape. "
                 "The scenario has no person and no place ('it "
                 "automatically detects and reminds') → send it back: 'Who? "
                 "Where? Say the name.' A student starts agonizing over "
                 "colors → 'Colors go on the later list. Today we decide "
                 "what it does.' Output anchor 6: the usage scenario. "
                 "Stand-still. The prompt box is 6.3 verbatim — shared text; "
                 "any edit must sync all three documents.")

    # ---------------------------------------------------------------- 14 requirements sheet
    s = page(prs, 14)
    add_title_bar(s, "This Table Is Your Project's Building Permit")
    add_multiline(s, 0.62, 1.00, 8.76, 0.52, [
        ("Its proper name is the requirements sheet — on a real job site, "
         "they call it requirements.md", {"size": 11}),
        ("— and it means one thing: BEFORE YOU START, WRITE DOWN EXACTLY "
         "WHAT \"IT DOES\" MEANS.", {"size": 11, "bold": True}),
    ], line_spacing=1.2)
    grey_box(s, 0.62, 1.60, 5.30, 3.60, [
        "# My Project Requirements Sheet",
        "Name: ______  Date: ______",
        "## Who it helps and what problem it solves",
        "   (one sentence, with a real name)",
        "## Usage scenario (who uses it, where, how — 4 sentences max)",
        "## Input (what it senses, 3 lines max)",
        "## Logic (what it does under what condition, 3 lines max;",
        "   include one: what happens with no signal / can't recognize)",
        "## Output (how it shows itself, 3 lines max)",
        "## Core feature (exactly 1)",
        "## The \"later\" list (what got cut lives here, 1 item min)",
        "1. ______ (why it was cut: ______)",
        "## Hardware check (fill after the tour)",
        "Modules needed: ______  in the pool / substitute: ______",
        "signature: ______",
    ], size=9, line_spacing=1.25)
    add_text(s, 6.10, 1.62, 3.28, 0.24, "WHY WRITE IT?", size=10.5, bold=True)
    reasons = [
        ("1", "It keeps you from drifting — every marathon work session "
              "checks it first"),
        ("2", "It's alive — you can change it later, but CROSS IT OUT, "
              "DON'T ERASE IT, so everyone can see you changed your mind"),
        ("3", "Three sections only, THREE LINES MAX EACH — no essays. Can't "
              "fit it = not thought through"),
    ]
    ry = 1.92
    for num, txt in reasons:
        num_block(s, 6.10, ry, num, size=0.32)
        add_text(s, 6.56, ry - 0.02, 2.82, 0.66, txt, size=10,
                 line_spacing=1.2)
        ry += 0.78
    add_rect(s, 6.10, 4.32, 3.28, 0.88, fill=YELLOW, border=None)
    add_multiline(s, 6.26, 4.40, 2.96, 0.74, [
        ("AI may help you polish the wording — but the facts must come from "
         "your two conversations just now.", {"size": 9.5}),
        ("Copy what AI invented, and you've given your project a fake birth "
         "certificate.", {"size": 9.5, "bold": True}),
    ], line_spacing=1.2)
    set_notes(s, "Look at this table on screen — it's your big project's "
                 "building permit. Its proper name is the requirements sheet "
                 "— on a real job site, they call it requirements.md — and it "
                 "means one thing: before you start, write down exactly what "
                 "'it does' means. Why write it? Three reasons. One: the "
                 "marathon starts at Lesson 8 — every time you start work, "
                 "you check it first. It keeps you from drifting. Two: it's "
                 "alive — you can change it later, but cross it out, don't "
                 "erase it, so everyone can see you changed your mind. "
                 "Three: if you still can't say it after writing three "
                 "sections, you haven't thought it through. No essays — three "
                 "sections only, three lines max each. Fill it in — in your "
                 "workbook or notebook. AI may help you polish the wording, "
                 "but the facts must come from your two conversations just now "
                 "— copy what AI invented, and you've given your project a "
                 "fake birth certificate. Students fill the first six rows "
                 "(the hardware row stays blank until after the tour). "
                 "Written as an essay → TA C runs the template ruler and "
                 "sends it back: three sections, 3 lines max each — can't fit "
                 "= not thought through, go cut. A student has AI write the "
                 "whole thing → the citation check: 'The names and scenes in "
                 "here — yours, or invented?' Output anchor 7: the first six "
                 "rows. Stand-still. After they finish, cut to slide 15 for "
                 "the 2-minute Brandy contract talk, then return here. The "
                 "template is 6.5 verbatim — identical to the workbook; any "
                 "edit must sync all three documents.")

    # ---------------------------------------------------------------- 15 Brandy
    s = page(prs, 15)
    add_title_bar(s, "Look What a Professional Engineer Writes First")
    add_multiline(s, 0.62, 1.05, 8.76, 0.56, [
        ("Brandy — Seeed's application engineer, the one who built a voice "
         "keyboard with ZERO code.", {"size": 12}),
        ("Before she starts, she writes an input/output contract — four "
         "lines:", {"size": 12}),
    ], line_spacing=1.3)
    contract = ["INPUT", "CONSTRAINTS", "OUTPUT", "SUCCESS CRITERIA"]
    for i, c in enumerate(contract):
        cx = 0.62 + i * 2.24
        add_card(s, cx, 1.80, 2.00, 0.72, fill=WHITE, border=INK,
                 border_w=1.5)
        add_text(s, cx, 1.80, 2.00, 0.72, c, size=12, bold=True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < 3:
            add_text(s, cx + 2.00, 1.80, 0.24, 0.72, "→", size=15, bold=True,
                     color=YELLOW, align=PP_ALIGN.CENTER,
                     anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    compares = [
        "HER \"who she makes it for\" = YOUR REAL USER",
        "HER \"use case\" = YOUR USAGE SCENARIO",
        "HER \"input and output\" = YOUR INPUT · LOGIC · OUTPUT",
        "HER \"success criteria\" = YOUR CORE FEATURE + WHAT COUNTS AS "
        "SUCCESS",
    ]
    cy_ = 2.78
    for i, c in enumerate(compares):
        num_block(s, 0.62, cy_, str(i + 1), size=0.30)
        add_text(s, 1.06, cy_ - 0.01, 8.32, 0.28, c, size=11)
        cy_ += 0.38
    add_rect(s, 0.62, 4.38, 0.07, 0.84, fill=YELLOW, border=None)
    add_multiline(s, 0.88, 4.42, 8.50, 0.78, [
        ("SHE CALLS IT A CONTRACT; YOU CALL IT A REQUIREMENTS SHEET — IT'S "
         "THE SAME THING. You're already using an engineer's method.",
         {"size": 11.5, "bold": True}),
        ("Starting next session, every new module you take on can start with "
         "those four lines.", {"size": 10.5}),
    ], line_spacing=1.25)
    set_notes(s, "You just filled a requirements sheet. Now look at what a "
                 "professional engineer writes before starting — Brandy, "
                 "Seeed's application engineer, the one who built a voice "
                 "keyboard with zero code. Before she starts, she writes an "
                 "input/output contract: input, constraints, output, success "
                 "criteria — four lines. Compare with yours: your real user "
                 "= who she makes it for; your usage scenario = her use case; "
                 "your input·logic·output = her input and output; your core "
                 "feature + what counts as success = her success criteria. She "
                 "calls it a contract; you call it a requirements sheet — "
                 "it's the same thing. You're already using an engineer's "
                 "method. Starting next session, every new module you take on "
                 "can start with those four lines. This page is a V4 addition "
                 "the CN v2 plan never had — the 19th page. Cut in once after "
                 "the students finish the first six rows, then immediately "
                 "return to slide 14 (the stand-still base). No output anchor "
                 "— the talk is framing, not a task.")

    # ---------------------------------------------------------------- 16 scope killer (RED 2)
    s = page(prs, 16)
    add_title_bar(s, "The Scope Killer — Exactly One Core Feature")
    add_multiline(s, 0.62, 1.02, 8.76, 0.56, [
        ("Send AI your requirements sheet. It may ask ONLY ONE question: "
         "\"Would it die without this?\"", {"size": 12, "bold": True}),
        ("It may not say \"this feature is bad.\"", {"size": 11.5}),
    ], line_spacing=1.25)
    red_line(s, 1.66, "IT MAY NOT DECIDE FOR YOU. THE VETO IS ALWAYS YOURS.",
             size=14)
    grey_box(s, 0.62, 2.30, 8.76, 1.60, [
        "This is my requirements sheet: (paste the whole thing).",
        "Act as the \"Scope Killer\": for every feature I wrote, ask me only",
        "\"Would it die without this?\" — one at a time, one question each.",
        "Don't decide for me. I make the final call on what gets cut.",
        "Whatever gets cut, organize into a \"later\" list for me.",
    ], size=10, line_spacing=1.3)
    add_rect(s, 0.62, 4.02, 0.07, 0.66, fill=YELLOW, border=None)
    add_multiline(s, 0.88, 4.06, 8.50, 0.60, [
        ("The core-feature row holds exactly one item. Whatever gets cut goes "
         "into the \"later\" list, one line each, with the reason.",
         {"size": 10.5}),
        ("CUT THINGS KEEP THEIR NAMES — ONLY THEN HAS A TRADE-OFF ACTUALLY "
         "HAPPENED.", {"size": 10.5, "bold": True}),
    ], line_spacing=1.25)
    yellow_box(s, 0.62, 4.72, 8.76, 0.50, [
        "Core feature = exactly 1  |  \"later\" list 1 item min, each with "
        "its reason."], size=10)
    set_notes(s, "Last step — bring in the Scope Killer. The rules: send AI "
                 "your requirements sheet, and it may ask only one question: "
                 "'Would it die without this?' It may not say 'this feature is "
                 "bad' — the veto is always yours. Goal: the core feature row "
                 "holds exactly one item. Whatever gets cut goes into the "
                 "'later' list, one line each, with the reason. Cut things "
                 "keep their names — only then has a trade-off actually "
                 "happened. The red line is the deck's second red — the veto "
                 "ownership. Not a single cut → watch closely: it's not a "
                 "perfect scope, it's an unwillingness to let go. A TA sits "
                 "down and cuts one with them. Cut down to a hollow shell "
                 "(even the core got cut) → rescue it: 'Which function, if "
                 "gone, makes it not it anymore? That one comes back.' AI "
                 "itself starts inflating the scope ('should I add "
                 "networking?') → teaching point, say it to the whole class "
                 "once: 'See — AI can talk a need bigger and bigger too. It "
                 "needs managing as much as anything else. Interrupt it: "
                 "First ask me: would it die without this?' Output anchor 8: "
                 "the cut record. Stand-still. The prompt box is 6.4 "
                 "verbatim — shared text; any edit must sync all three "
                 "documents.")

    # ---------------------------------------------------------------- 17 parts tour (RED 3)
    s = page(prs, 17)
    add_title_bar(s, "The Parts-Pool Tour — Not Window-Shopping")
    add_multiline(s, 0.62, 1.00, 8.76, 0.50, [
        ("Requirements sheet in hand — one last question: can the things in "
         "your hands actually make this?", {"size": 11.5}),
        ("The parts pool has the answer.", {"size": 11.5}),
    ], line_spacing=1.25)
    tables = [
        ("Table 1", "SENSE (sensors/cameras)"),
        ("Table 2", "OUTPUT (lights/buzzers/screens)"),
        ("Table 3", "ACTUATION (servos/motors)"),
        ("Table 4", "MIXED"),
    ]
    for i, (t1, t2) in enumerate(tables):
        tx = 0.62 + i * 2.22
        add_rect(s, tx, 1.58, 2.10, 0.62, fill=YELLOW, border=None)
        add_multiline(s, tx + 0.10, 1.64, 1.90, 0.52, [
            (t1, {"size": 10.5, "bold": True}),
            (t2, {"size": 8.5}),
        ], line_spacing=1.15)
    add_card(s, 0.62, 2.36, 8.76, 1.86, fill=WHITE, border=INK, border_w=1.5)
    add_multiline(s, 0.84, 2.46, 8.32, 1.68, [
        ("1. Modules my sheet needs: ______ → found the real thing? (tick)",
         {"size": 10.5, "mono": True}),
        ("2. Not in the pool: ______ → substitute: ______ (ask a TA, check "
         "the four-table sort)", {"size": 10.5, "mono": True}),
        ("3. Discovery picks: touch one module you've never seen at each "
         "table, write its name + your guess", {"size": 10.5, "mono": True}),
        ("Table 1: ______   Table 2: ______   Table 3: ______",
         {"size": 10.5, "mono": True}),
        ("4. Back at your seat: fill the \"hardware check\" row and sign",
         {"size": 10.5, "mono": True}),
    ], line_spacing=1.35)
    add_text(s, 0.62, 4.32, 8.76, 0.26,
             "Four groups rotate; 7 minutes per table; the whistle moves you.",
             size=11, bold=True)
    add_rect(s, 0.62, 4.66, 0.07, 0.54, fill=RED, border=None)
    add_oval(s, 0.86, 4.72, 0.12, fill=RED, border=None)
    add_multiline(s, 1.12, 4.68, 8.26, 0.50, [
        ("NOT IN THE POOL? FIND A SUBSTITUTE — OR IT GOES ON THE \"LATER\" "
         "LIST.", {"size": 10, "bold": True, "color": RED}),
        ("CHANGE THE NEED, NOT THE PURCHASE — today is look, touch, and "
         "register only; checking parts out happens at the Lesson 8 "
         "marathon start.", {"size": 10, "bold": True, "color": RED}),
    ], line_spacing=1.2)
    set_notes(s, "Requirements sheet in hand — one last question: can the "
                 "things in your hands actually make this? — The parts pool "
                 "has the answer. The tour is not window-shopping. Four "
                 "tasks (project them): 1. tick the modules your sheet needs "
                 "and find the physical parts; 2. not in the pool — find a "
                 "substitute using the four-table sort on screen; 3. the "
                 "three discovery picks — touch one module you've never seen "
                 "at each table, guess what it does, check with a TA; 4. back "
                 "at your seat, sign the hardware row. Four groups rotate; 7 "
                 "minutes per table; the whistle moves you. The red line is "
                 "the deck's third red — the iron line: change the need, not "
                 "the purchase. The tour turns into a general-store stroll → "
                 "the task sheet is the brake: TAs ask only one question per "
                 "table — 'Which task are you on?' The module they need isn't "
                 "in the pool → 'An engineer's first lesson: solve the "
                 "problem with what you have. Can't buy it? It goes on the "
                 "later list.' (we don't promise purchases). A student wants "
                 "to take a module away → make it clear: today is look, touch, "
                 "and register only; checking parts out happens at the "
                 "Lesson 8 marathon start. 4 groups × 4 tables (7 min each + "
                 "1 min rotation); complete the task sheet; sign the hardware "
                 "row. Output anchor 9: the hardware row signed. The task "
                 "sheet is 6.6 verbatim — identical to the workbook; any edit "
                 "must sync all three documents. (Buffer 16:45–16:50: no "
                 "slide — this page stays up or the screen goes dark; the "
                 "instructor counts acceptance numbers and gets the Brief "
                 "Wall ready.)")

    # ---------------------------------------------------------------- 18 ceremony
    s = page(prs, 18)
    add_text(s, 0.62, 1.30, 8.76, 0.60, "MY PROJECT GOT ITS BRIEF TODAY.",
             size=30, bold=True, align=PP_ALIGN.CENTER)
    actions = [
        "Take a sticky note: project name + core feature in one sentence",
        "Line up — stick it on the Brief Wall",
        "Say it aloud as you stick it",
    ]
    for i, a in enumerate(actions):
        ax = 0.62 + i * 2.96
        add_card(s, ax, 2.20, 2.84, 0.90, fill=WHITE, border=INK,
                 border_w=1.5)
        num_block(s, ax + 0.14, 2.32, str(i + 1), size=0.32)
        add_text(s, ax + 0.60, 2.30, 2.12, 0.72, a, size=9.5,
                 line_spacing=1.2)
    add_rect(s, 0.62, 3.40, 0.07, 1.30, fill=YELLOW, border=None)
    add_multiline(s, 0.88, 3.44, 8.50, 1.24, [
        ("The requirements sheet stays in your workbook — it's your building "
         "permit; every marathon work session starts by checking it.",
         {"size": 11.5}),
        ("The stickies on this wall are the whole class's reference. It can "
         "change — but CROSS IT OUT, DON'T ERASE IT, so everyone can see you "
         "changed your mind.", {"size": 11.5, "bold": True}),
    ], line_spacing=1.35)
    set_notes(s, "Everyone take a sticky note. Write two things: your "
                 "project's name, and its core feature in one sentence. Line "
                 "up — and stick it on the Brief Wall. Say it aloud as you "
                 "stick it. (one by one, stick and say) The requirements "
                 "sheet itself stays in your workbook — it's your building "
                 "permit; every marathon work session starts by checking it. "
                 "The stickies on this wall are the whole class's reference. "
                 "It can change — but cross it out, don't erase it, so "
                 "everyone can see you changed your mind. A student's sheet "
                 "isn't finished → the sticky still goes up (a name + core "
                 "feature is enough); the sheet is completed after class. "
                 "The ceremony doesn't wait. Can't say the core feature "
                 "aloud → the neighbor reads it and the student nods — no "
                 "skipping; named participation is part of the ceremony. This "
                 "page is the lesson's ceremony peak — deliberately bare, "
                 "deliberately set apart from the operation pages (same "
                 "register as slide 08). The Brief Wall is the physical wall "
                 "next to the Project Wall (set up before class).")

    # ---------------------------------------------------------------- 19 close
    s = page(prs, 19)
    add_title_bar(s, "Today, You Really Did a Trade-Off")
    review = [
        "You put your project through the wringer — AI opened the list, your "
        "classmate tested it for real",
        "You made a call on every fault — accept, or reject — and you can "
        "say why",
        "Your topic walked off a wall and into a requirements sheet — you "
        "cut things, and you can say why",
    ]
    ry = 1.05
    for r in review:
        check_item(s, 0.62, ry, 8.76, r, size=11, h=0.30)
        ry += 0.38
    add_card(s, 0.62, 2.24, 8.76, 1.00, fill=WHITE, border=INK, border_w=1.5)
    add_multiline(s, 0.84, 2.34, 8.32, 0.82, [
        ("Four lines as usual — today's \"done\" line:",
         {"size": 11, "bold": True}),
        ("MY BIG PROJECT GOT ITS BRIEF TODAY — IT'S CALLED ___, ITS CORE "
         "FEATURE IS ___, I CUT ___", {"size": 12.5, "bold": True}),
    ], line_spacing=1.4)
    yellow_box(s, 0.62, 3.38, 8.76, 0.92, [
        "Fab Academy graduates say that on the day their Final Project "
        "finished, they realized: it wasn't just a graduation project — it "
        "was THE START OF A LIFE PROJECT. The sentence you wrote today is "
        "that too."], size=11)
    add_rect(s, 0.62, 4.42, 8.76, 0.62, fill=WHITE, border=None)
    add_multiline(s, 0.80, 4.48, 8.40, 0.52, [
        ("NEXT TIME — the marathon's first leg. Bring your requirements "
         "sheet: check out parts, start building, the stand-up — and AI as "
         "reviewer appears for the first time. Today you said \"I want to "
         "build.\" Next session, you'll say \"I'm building.\"", {"size": 10}),
    ], line_spacing=1.2)
    add_text(s, 0.62, 5.06, 8.76, 0.20,
             "Sticky notes on the wall stay. Boards in the box.", size=9.5)
    set_notes(s, "Fab Academy graduates say that on the day their Final "
                 "Project finished, they realized: it wasn't just a "
                 "graduation project — it was the start of a life project. "
                 "The sentence you wrote today is that too. Look back at "
                 "today: this morning you put your project through the "
                 "wringer and made Accept-or-Reject calls. This afternoon, "
                 "your topic walked off a wall and into a requirements sheet "
                 "— you cut things, and you can say why. That's 'making "
                 "trade-offs' — and today, you really did it. Log, four lines "
                 "as usual; 'done' reads: my big project got its brief today "
                 "— it's called ___, its core feature is ___, I cut ___. Next "
                 "session preview: the marathon's first leg. Bring your "
                 "requirements sheet: check out parts, start building, the "
                 "stand-up — and AI as reviewer appears for the first time. "
                 "Your project already has a home on your own computer; the "
                 "brief is already on the wall — everything's ready. Run. "
                 "Today you said 'I want to build.' Next session, you'll say "
                 "'I'm building.' No ability-card ceremony today; the "
                 "ceremony peak was slide 18, and this close settles the arc "
                 "+ previews the marathon.")

    prs.save(OUT)
    print(f"Saved: {OUT} slides: {len(prs.slides._sldIdLst)}")


if __name__ == "__main__":
    main()
