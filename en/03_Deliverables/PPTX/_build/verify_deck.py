"""Verify a built EN deck against the Chaihuo brand rules:
1) no CJK characters in any on-screen text (all English/digits)
2) footer page numbers are NN / TOTAL and match slide order
3) red element count per page (deck budget: controlled)
4) no shape runs outside the 10 x 5.625 in canvas
5) slide count matches expected total
"""
import sys, re, os
from pptx import Presentation
from pptx.util import Emu

def cjk_scan(text):
    return re.findall(r'[\u4e00-\u9fff\u3000-\u303f\uff00-\uffef]', text)

def main(path, expected_total):
    prs = Presentation(path)
    W, H = prs.slide_width, prs.slide_height
    n = len(prs.slides._sldIdLst)
    print(f"file: {os.path.basename(path)}")
    print(f"slides: {n} (expected {expected_total})  canvas: {W/914400:.3f}x{H/914400:.3f} in")
    ok = True
    if n != expected_total:
        ok = False
        print("  FAIL slide count")
    for idx, slide in enumerate(prs.slides, start=1):
        texts = []
        reds = 0
        for sh in slide.shapes:
            # geometry check
            try:
                x, y, w, h = sh.left, sh.top, sh.width, sh.height
                if x is not None and y is not None and w is not None and h is not None:
                    if x < -1000 or y < -1000 or x + w > W + 1000 or y + h > H + 1000:
                        print(f"  slide {idx:02d}: shape out of canvas? x={x/914400:.2f} y={y/914400:.2f} "
                              f"w={w/914400:.2f} h={h/914400:.2f}")
                        ok = False
            except Exception:
                pass
            if sh.shape_type == 13:  # picture
                continue
            if sh.has_text_frame:
                t = sh.text_frame.text
                if t.strip():
                    texts.append(t)
            # red fill count
            try:
                if sh.fill.type is not None and str(sh.fill.type) == "MSO_FILL_TYPE.SOLID (1)":
                    if sh.fill.fore_color.rgb is not None and str(sh.fill.fore_color.rgb) == "D84144":
                        reds += 1
            except Exception:
                pass
        full = "\n".join(texts)
        cjk = cjk_scan(full)
        if cjk:
            print(f"  slide {idx:02d}: CJK FOUND: {sorted(set(cjk))}")
            ok = False
        # footer check
        footers = [t for t in texts if "|" in t and re.search(r'\d{2} / \d{2}', t)]
        if footers:
            for f in footers:
                m = re.search(r'\| (\d{2}) / (\d{2})', f)
                if m and int(m.group(1)) != idx:
                    print(f"  slide {idx:02d}: footer mismatch: '{f}'")
                    ok = False
        else:
            # cover may omit footer? our builder adds footer to all pages
            print(f"  slide {idx:02d}: NO footer found")
            ok = False
        if reds:
            print(f"  slide {idx:02d}: red elements = {reds}")
    print("RESULT:", "PASS" if ok else "FAIL")

if __name__ == "__main__":
    main(sys.argv[1], int(sys.argv[2]))
